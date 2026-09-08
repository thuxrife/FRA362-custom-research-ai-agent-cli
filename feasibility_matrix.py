from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_telos_sdg_excel(topic: str, out_path: Path, custom_scores: dict = None) -> Path:
    wb = openpyxl.Workbook()
    
    # Visual Styles
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    category_header_font = Font(name="Segoe UI", size=11, bold=True, color="0F172A")
    title_font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    bold_font = Font(name="Segoe UI", size=10, bold=True, color="0F172A")
    regular_font = Font(name="Segoe UI", size=9, color="1E293B")
    verdict_font = Font(name="Segoe UI", size=11, bold=True, color="1E3A8A")
    
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    cat_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    total_fill = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
    rubric_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # -------------------------------------------------------------------------
    # Sheet 1: Detailed TELOS-SDG Matrix (Multi-Question per Topic)
    # -------------------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "TELOS_SDG_Matrix"
    
    ws1["A1"] = "TELOS-SDG Multi-Criteria Feasibility Matrix"
    ws1["A1"].font = title_font
    ws1["A2"] = f"Evaluation Target: {topic} | Predetermined Goal-First Scoring (Anti-Backpropagation)"
    ws1["A2"].font = subtitle_font
    
    headers = ["ID", "Pillar & Diagnostic Question", "Weight (%)", "Score (1-5)", "Weighted Score (0-100)", "Consulting Evidence & Rationale"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    # Detailed sub-questions per topic
    questions_data = [
        # Technology (Total 20%)
        ("CAT_TECH", "1. TECHNOLOGY FEASIBILITY (Weight: 20%)", "", "", "", "Audit: Team Capability vs. Novelty Risk"),
        ("T1.1", "Can our actual human team execute this with current skills? (teammate-persona audit)", 7.0, 4.0, "Core Python/backend skills present; minimal upskilling required for APIs."),
        ("T1.2", "Is the technology proven in the industry (Pioneer Red Flag Check)?", 7.0, 4.5, "CLEARED: Established commercial precedents exist; we are NOT unproven pioneers."),
        ("T1.3", "Are external hardware/cloud APIs and third-party dependencies production-stable?", 6.0, 4.0, "Mature cloud endpoints with 99.9% SLA availability."),

        # Economic (Total 15%)
        ("CAT_ECON", "2. ECONOMIC FEASIBILITY (Weight: 15%)", "", "", "", "Audit: Value vs. Financial Constraints"),
        ("E2.1", "Is the return worth the capital, opportunity cost, and financial risk?", 5.0, 3.5, "Strong ROI justified by mitigation of recurring disaster losses."),
        ("E2.2", "Can upfront CapEx and recurring OpEx comfortably stay within budget/runway?", 5.0, 4.0, "Low infrastructure run-rate using serverless cloud microservices."),
        ("E2.3", "Is there a real-world benchmark proving financial attractiveness / payback?", 5.0, 3.5, "Comparable municipal deployments report payback within 8-10 months."),

        # Legal (Total 15%)
        ("CAT_LEGAL", "3. LEGAL & SOCIAL FEASIBILITY (Weight: 15%)", "", "", "", "Audit: Law, Social Norms & Pending Bills"),
        ("L3.1", "Does the solution comply with statutory laws (PDPA, GDPR, IP, Municipal Code)?", 5.0, 4.5, "Full compliance with local privacy acts; no proprietary IP infringement."),
        ("L3.2", "Does the project align with social norms, ethics, and public acceptance (zero boycott risk)?", 5.0, 4.0, "High community trust; transparent public alert protocols prevent backlash."),
        ("L3.3", "What is the risk exposure to pending legislation or emerging regulatory mandates?", 5.0, 3.5, "Monitored emerging municipal AI governance mandates; compliance hooks included."),

        # Operational (Total 20% - Human/Team Centric)
        ("CAT_OPER", "4. OPERATIONAL FEASIBILITY (Weight: 20% - Human & Team Focus)", "", "", "", "Audit: Human Friction Before vs. After Launch"),
        ("O4.1", "Before Launch: What human changes, new roles, or process overhauls are required?", 5.0, 3.0, "Requires 2 weeks of operational workflow training for ground crew."),
        ("O4.2", "Before Launch: Are critical datasets and tools accessible without bureaucratic silos?", 5.0, 3.5, "Data pipeline requires inter-departmental security credentials."),
        ("O4.3", "After Launch: How severely will ongoing operations divert senior team bandwidth?", 5.0, 3.0, "Senior bandwidth impact estimated at 15% during initial 4-week rollout."),
        ("O4.4", "After Launch: Who handles on-call 2 AM escalation and prevents operator burnout?", 5.0, 3.5, "Shared rotational on-call schedule prevents single-point-of-failure burnout."),

        # Schedule (Total 15%)
        ("CAT_SCHED", "5. SCHEDULE FEASIBILITY (Weight: 15%)", "", "", "", "Audit: Critical Path & Delivery Guarantees"),
        ("S5.1", "Is the deadline realistically achievable given scope and team capacity?", 5.0, 3.5, "Feasible for 12-week MVP if discovery phase finishes on time."),
        ("S5.2", "What non-negotiable conditions (frozen scope, early access) guarantee delivery?", 5.0, 4.0, "Strict scope lock by Week 2 avoids mid-sprint scope creep."),
        ("S5.3", "Is there an explicit MVP de-scoping plan if critical-path delays occur?", 5.0, 3.5, "Non-essential analytics dashboard flagged for de-scoping if Week 8 slips."),

        # SDGs (Total 15% - Wedding Cake Model)
        ("CAT_SDG", "6. UN SDGs WEDDING CAKE FEASIBILITY (Weight: 15%)", "", "", "", "Audit: Multi-Tier Targets (Biosphere, Society, Economy)"),
        ("SDG6.1", "Does the project measurably contribute to specific UN SDG Targets (e.g. 6.3, 11.5)?", 5.0, 4.5, "Directly advances Target 11.5 (Disaster Reduction) and Target 6.3 (Water Quality)."),
        ("SDG6.2", "Does the solution span across >1 layer of the SDG Wedding Cake?", 5.0, 5.0, "CONFIRMED: Anchored in Biosphere (Water), Society (Urban), and Economy (Assets)."),
        ("SDG6.3", "Does the solution avoid unintended negative trade-offs across other SDG tiers?", 5.0, 4.0, "Safe ecological discharge parameters prevent downstream environmental harm.")
    ]

    current_row = 5
    question_rows = []
    
    for item in questions_data:
        q_id, title, weight, default_score, notes = item[0], item[1], item[2], item[3], item[4]
        
        # Category Heading Row
        if q_id.startswith("CAT_"):
            c1 = ws1.cell(row=current_row, column=1, value="")
            c2 = ws1.cell(row=current_row, column=2, value=title)
            c3 = ws1.cell(row=current_row, column=3, value="")
            c4 = ws1.cell(row=current_row, column=4, value="")
            c5 = ws1.cell(row=current_row, column=5, value="")
            c6 = ws1.cell(row=current_row, column=6, value=notes)
            
            c2.font = category_header_font
            c6.font = subtitle_font
            for c in range(1, 7):
                ws1.cell(row=current_row, column=c).fill = cat_fill
                ws1.cell(row=current_row, column=c).border = thin_border
            current_row += 1
            continue
            
        # Question Row
        question_rows.append(current_row)
        score_val = custom_scores.get(q_id, default_score) if custom_scores else default_score
        
        c_id = ws1.cell(row=current_row, column=1, value=q_id)
        c_title = ws1.cell(row=current_row, column=2, value=title)
        c_weight = ws1.cell(row=current_row, column=3, value=weight / 100.0)
        c_score = ws1.cell(row=current_row, column=4, value=score_val)
        c_wscore = ws1.cell(row=current_row, column=5, value=f"=(D{current_row}*20)*(C{current_row})") # (Score * 20) * Weight%
        c_notes = ws1.cell(row=current_row, column=6, value=notes)
        
        c_id.font = bold_font
        c_title.font = regular_font
        c_weight.font = regular_font
        c_score.font = bold_font
        c_wscore.font = bold_font
        c_notes.font = regular_font
        
        c_weight.number_format = "0.0%"
        c_score.number_format = "0.0"
        c_wscore.number_format = "0.0"
        
        c_id.alignment = Alignment(horizontal="center", vertical="center")
        c_weight.alignment = Alignment(horizontal="center", vertical="center")
        c_score.alignment = Alignment(horizontal="center", vertical="center")
        c_wscore.alignment = Alignment(horizontal="center", vertical="center")
        c_title.alignment = Alignment(wrap_text=True, vertical="center")
        c_notes.alignment = Alignment(wrap_text=True, vertical="center")
        
        for c in range(1, 7):
            ws1.cell(row=current_row, column=c).border = thin_border

        current_row += 1

    # Total Summary Row
    total_row = current_row
    ws1.cell(row=total_row, column=1, value="TOTAL").font = bold_font
    ws1.cell(row=total_row, column=2, value="Cumulative Feasibility Score & Decision Verdict").font = bold_font
    
    # Sum weights and sum weighted scores
    weight_formula = "=SUM(" + ",".join([f"C{r}" for r in question_rows]) + ")"
    score_formula = "=SUM(" + ",".join([f"E{r}" for r in question_rows]) + ")"
    verdict_formula = f'=IF(E{total_row}>=80, "HIGHLY VIABLE - GREEN LIGHT", IF(E{total_row}>=60, "CONDITIONALLY VIABLE - PROCEED WITH MITIGATION", "HIGH RISK / UNVIABLE"))'
    
    ws1.cell(row=total_row, column=3, value=weight_formula).number_format = "0.0%"
    ws1.cell(row=total_row, column=5, value=score_formula).number_format = "0.0"
    ws1.cell(row=total_row, column=6, value=verdict_formula).font = verdict_font

    for c in range(1, 7):
        cell = ws1.cell(row=total_row, column=c)
        cell.fill = total_fill
        cell.border = thin_border
        if c in [1, 2, 3, 5]:
            cell.font = bold_font

    ws1.column_dimensions['A'].width = 12
    ws1.column_dimensions['B'].width = 50
    ws1.column_dimensions['C'].width = 14
    ws1.column_dimensions['D'].width = 14
    ws1.column_dimensions['E'].width = 20
    ws1.column_dimensions['F'].width = 55

    # -------------------------------------------------------------------------
    # Sheet 2: Clear Scoring Rubric & Standard (1-5 Scale Definition)
    # -------------------------------------------------------------------------
    ws_rubric = wb.create_sheet(title="Scoring_Rubric_Guide")
    ws_rubric["A1"] = "TELOS-SDG Standardized Scoring System & Rubric"
    ws_rubric["A1"].font = title_font
    ws_rubric["A2"] = "Clear, non-subjective evaluation standards for scoring every diagnostic question (1-5 scale)."
    ws_rubric["A2"].font = subtitle_font
    
    rubric_headers = ["Score", "Rating Level", "General Meaning", "Technology Criterion", "Economic Criterion", "Operational Criterion (Team)"]
    for c_idx, h in enumerate(rubric_headers, 1):
        cell = ws_rubric.cell(row=4, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    rubric_rows = [
        (1, "Critical Barrier / Unviable", "Showstopper roadblock. Severe risk, lack of baseline capacity, or massive harm.", "No team capability; unproven technology; pioneer red flag triggered.", "Severe negative cash flow; payback > 5 years; cost exceeds total runway.", "Massive team burnout; complete lack of access to critical data silos."),
        (2, "High Risk / Significant Gaps", "Major deficits. Requires substantial external hiring or heavy compromise.", "Team lacks core stack skills; requires extensive hiring or custom R&D.", "Marginal ROI; high hidden operational costs; significant financial risk.", "Heavy friction; senior staff diverted >30%; severe resistance to change."),
        (3, "Moderate / Conditionally Viable", "Standard manageable difficulty. Feasible with clear mitigation plan.", "Team has related skills; standard learning curve; proven open libraries.", "Acceptable ROI; CapEx within budget; payback horizon within 12-18 months.", "Manageable friction; training needed for 2 weeks; clear on-call rotation."),
        (4, "Strong / High Viability", "Strong alignment. Team possesses skills; proven commercial precedents.", "Existing team mastery; mature cloud endpoints; clear architectural blueprint.", "Strong value proposition; clear payback within 6-12 months; minimal CapEx.", "Smooth integration; minimal disruption to ongoing works; enthusiastic team buy-in."),
        (5, "Optimal / Outstanding", "Best-in-class. Immediate unfair advantage; zero friction; systemic impact.", "Elite team mastery; established open standard; zero pioneer novelty risk.", "Exceptional ROI; self-sustaining economics; immediate cost reduction.", "Empowers human team; streamlines daily operations; zero on-call burnout.")
    ]
    
    for r_idx, r_data in enumerate(rubric_rows, start=5):
        for c_idx, val in enumerate(r_data, start=1):
            cell = ws_rubric.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            cell.font = bold_font if c_idx in [1, 2] else regular_font
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(wrap_text=True, vertical="center")
                
    ws_rubric.column_dimensions['A'].width = 10
    ws_rubric.column_dimensions['B'].width = 25
    ws_rubric.column_dimensions['C'].width = 30
    ws_rubric.column_dimensions['D'].width = 35
    ws_rubric.column_dimensions['E'].width = 35
    ws_rubric.column_dimensions['F'].width = 35

    # -------------------------------------------------------------------------
    # Sheet 3: Goals & Clear Direction (Anti-Backpropagation)
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Goals_and_Direction")
    ws3["A1"] = "Pre-Established Feasibility Goals & Idea Direction"
    ws3["A1"].font = title_font
    ws3["A2"] = "Rule: Ideas are scored against these predetermined goals to prevent score back-propagation."
    ws3["A2"].font = subtitle_font
    
    goals_rows = [
        ("1. Primary Project Goal", f"Deliver actionable, high-impact outcome for: {topic}"),
        ("2. Solution Direction", "Autonomous, human-centric, and sustainable deployment within constraints."),
        ("3. Target Beneficiaries", "Direct users, operations team, and community stakeholders."),
        ("4. Non-Negotiable Constraints", "Budget cap, team availability, zero legal violation, and positive SDG alignment.")
    ]
    
    for r_idx, (g_title, g_desc) in enumerate(goals_rows, start=4):
        ws3.cell(row=r_idx, column=1, value=g_title).font = bold_font
        ws3.cell(row=r_idx, column=2, value=g_desc).font = regular_font
        ws3.cell(row=r_idx, column=1).border = thin_border
        ws3.cell(row=r_idx, column=2).border = thin_border

    ws3.column_dimensions['A'].width = 30
    ws3.column_dimensions['B'].width = 70

    # -------------------------------------------------------------------------
    # Sheet 4: SDG Wedding Cake Multi-Tier Mapping
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="SDG_Wedding_Cake")
    ws4["A1"] = "SDG Wedding Cake Multi-Tier Assessment"
    ws4["A1"].font = title_font
    ws4["A2"] = "Stockholm Resilience Centre Model: Biosphere -> Society -> Economy (Goal 17 Partnership)"
    ws4["A2"].font = subtitle_font
    
    cake_headers = ["Wedding Cake Layer", "Targeted SDG Goals", "Specific UN Target", "Official UN Indicator", "Feasibility Diagnostic Question"]
    for c_idx, h in enumerate(cake_headers, 1):
        cell = ws4.cell(row=4, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        
    cake_tiers = [
        ("Tier 1: Biosphere (Foundation)", "SDG 6 (Water), 13 (Climate)", "Target 6.3 & 13.1", "Indicator 6.3.2 & 13.1.2", "Does the system monitor/prevent toxic runoff and enhance municipal climate hazard resilience?"),
        ("Tier 2: Society (Middle)", "SDG 11 (Cities), 3 (Health)", "Target 11.5 & 3.9", "Indicator 11.5.1, 11.5.2 & 3.9.2", "Does the solution quantify direct reduction of economic asset losses and eliminate waterborne disease risks?"),
        ("Tier 3: Economy (Top)", "SDG 9 (Infra), 8 (Work), 12 (Prod)", "Target 9.1 & 12.5", "Indicator 9.1.1 & 12.5.1", "Does the infrastructure maintain >99.5% uptime while minimizing electronic e-waste through modular circular repair?"),
        ("Cross-Cutting: Partnerships", "SDG 17 (Partnerships)", "Target 17.18", "Indicator 17.18.1", "Does the project generate open, standardized telemetry streams feeding directly into municipal and SDG monitoring?")
    ]
    
    for r_idx, (layer, goals, targets, indicators, diag_q) in enumerate(cake_tiers, start=5):
        ws4.cell(row=r_idx, column=1, value=layer).font = bold_font
        ws4.cell(row=r_idx, column=2, value=goals).font = regular_font
        ws4.cell(row=r_idx, column=3, value=targets).font = bold_font
        ws4.cell(row=r_idx, column=4, value=indicators).font = regular_font
        ws4.cell(row=r_idx, column=5, value=diag_q).font = regular_font
        for col_i in range(1, 6):
            ws4.cell(row=r_idx, column=col_i).border = thin_border
            ws4.cell(row=r_idx, column=col_i).alignment = Alignment(wrap_text=True, vertical="center")

    ws4.column_dimensions['A'].width = 28
    ws4.column_dimensions['B'].width = 30
    ws4.column_dimensions['C'].width = 22
    ws4.column_dimensions['D'].width = 28
    ws4.column_dimensions['E'].width = 45

    wb.save(out_path)
    return out_path

if __name__ == '__main__':
    p = Path("test_matrix_v2.xlsx")
    generate_telos_sdg_excel("Bangkok Flood Mitigation", p)
    print("Multi-question matrix generated successfully! Size:", p.stat().st_size)
    p.unlink()
