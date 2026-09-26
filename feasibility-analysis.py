#!/usr/bin/env python3
"""
feasibility-analysis.py
-----------------------
เครื่องมือแปลงข้อมูลการประเมินของ AI Agents เป็น Excel Matrix (TELOS+S Horizontal Compiler)
หน้าที่ของสคริปต์นี้:
  - รับข้อมูลการประเมินจาก AI Agents (Group 2: The Jury และผู้เชี่ยวชาญ 6 ด้าน) ในรูปแบบ JSON
  - ไม่มี (Zero) การล็อกหรือเขียนคำถามไว้ในโค้ด Python แม้แต่ข้อเดียว
  - ตรวจสอบและบังคับคะแนนเป็นจำนวนเต็มเด็ดขาด {1, 2, 3, 4, 5} เท่านั้น
  - จัดรูปแบบเป็น Transposed Horizontal Matrix ภาษาไทย (1 Solution ต่อ 1 Sheet)
  - สร้างหน้าสรุปเปรียบเทียบทุก Solution (Executive Comparison Tab)
  - บันทึกไฟล์ลงใน 'feasibility-outcome/{month}-{date}-{year}-{time}_{seq}.xlsx'

การใช้งาน:
    python feasibility-analysis.py --input-json feasibility-outcome/jury_eval_data.json
    python feasibility-analysis.py (จะโหลด feasibility-outcome/jury_eval_data.json โดยอัตโนมัติ)
"""

import os
import sys
import json
import argparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string

# Ensure UTF-8 printing in Windows terminals
if sys.stdout and hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if sys.stderr and hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')


# ==============================================================================
# Helper: Canonical Pillar Normalization & Weighting
# ==============================================================================

def get_canonical_pillar(pillar_str, has_sdg=False):
    """
    จัดกลุ่มข้อความมิติการประเมินให้เป็น Canonical Key และค่าน้ำหนักมาตรฐาน (Two-Tier Model)
    รองรับทั้งภาษาไทย ภาษาอังกฤษ และรูปแบบข้อความที่ไม่สม่ำเสมอ
    """
    p = str(pillar_str).lower()
    if any(k in p for k in ["tech", "เทคนิค", "เทคโน"]):
        return ("T", 0.20)
    elif any(k in p for k in ["econ", "เศรษฐ", "การเงิน", "ต้นทุน"]):
        return ("E", 0.20)
    elif any(k in p for k in ["legal", "กฎหมาย", "ระเบียบ", "สถาบัน", "outreach"]):
        return ("L", 0.15 if has_sdg else 0.20)
    elif any(k in p for k in ["operat", "ปฏิบัติการ", "หน้างาน", "ผู้ใช้", "developer", "user"]):
        return ("O", 0.20)
    elif any(k in p for k in ["sched", "เวลา", "กำหนด", "แผนงาน"]):
        return ("S", 0.15 if has_sdg else 0.20)
    elif any(k in p for k in ["sdg", "ยั่งยืน", "สิ่งแวดล้อม"]):
        return ("SDG", 0.10)
    return ("OTHER", 0.20)


def make_pillar_sum_formula(col_letters, row_idx=11):
    """
    สร้างสูตร Excel สำหรับรวมคะแนนถ่วงน้ำหนักของคำถามในมิตินั้นๆ แล้วคูณด้วย 20
    เพื่อแปลงคะแนนสเกลถ่วงน้ำหนัก (เต็มตามค่าน้ำหนักมิติ เช่น 0.20) ให้เป็นคะแนนเต็มมิติ (เช่น 20.0)
    """
    if not col_letters:
        return "=0.0"
    col_indices = [column_index_from_string(c) for c in col_letters]
    # ตรวจสอบว่าเป็นคอลัมน์เรียงติดกันหรือไม่
    if len(col_indices) > 1 and col_indices == list(range(col_indices[0], col_indices[0] + len(col_indices))):
        return f"=SUM({col_letters[0]}{row_idx}:{col_letters[-1]}{row_idx})*20"
    elif len(col_indices) == 1:
        return f"={col_letters[0]}{row_idx}*20"
    else:
        cells = ",".join(f"{c}{row_idx}" for c in col_letters)
        return f"=SUM({cells})*20"


# ==============================================================================
# 1. ฟังก์ชันจัดรูปแบบและเขียนแผ่นงานสำหรับแต่ละ Solution (Transposed Horizontal Matrix)
# ==============================================================================

def write_solution_sheet(wb, sheet_title, assessment_data, solution_name="Solution", robotic_data=None):
    if not assessment_data:
        print(f"[แจ้งเตือน] {sheet_title} ไม่มีข้อมูล assessment_data ข้ามการสร้าง Sheet")
        return None

    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True

    # ฟอนต์และสี
    title_font = Font(name="Segoe UI", size=13, bold=True, color="1F497D")
    sub_font = Font(name="Segoe UI", size=9.5, italic=True, color="595959")
    attr_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    q_id_font = Font(name="Segoe UI", size=11, bold=True, color="1F497D")
    score_font = Font(name="Segoe UI", size=15, bold=True, color="002060")
    bold_font = Font(name="Segoe UI", size=9.5, bold=True)
    regular_font = Font(name="Segoe UI", size=9)
    verdict_font = Font(name="Segoe UI", size=12, bold=True, color="006100")

    attr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    score_hdr_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    score_cell_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
    highlight_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    q_id_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # หัวเรื่องตาราง
    ws.cell(row=1, column=1, value=f"เมทริกซ์ประเมินความเป็นไปได้เชิงระบบ TELOS+S: {solution_name}").font = title_font
    ws.cell(row=2, column=1, value="ระบบคะแนนจำนวนเต็มเด็ดขาด: 1, 2, 3, 4, 5 เท่านั้น (ไม่มีจุดทศนิยม / ไม่มี 0.5) | โครงสร้างแบบ Transposed แนวนอน").font = sub_font

    # นิยามแถวคุณลักษณะใน Column A
    attributes = [
        (4, "รหัสคำถาม (Question Code)", q_id_font, q_id_fill),
        (5, "มิติการประเมิน (TELOS+S Pillar)", attr_header_font, attr_fill),
        (6, "คำถามทดสอบความเสี่ยงจริง (Stress-Test Question)", attr_header_font, attr_fill),
        (7, "เหตุผลในการถาม & โหมดความล้มเหลวที่เฝ้าระวัง (Rationale)", attr_header_font, attr_fill),
        (8, "เกณฑ์การให้คะแนน (Scoring Rubric Definition: ระดับ 1, 2, 3, 4, 5 ครบทุกระดับ)", attr_header_font, attr_fill),
        (9, "คะแนนที่ได้รับ (ASSIGNED SCORE: 1, 2, 3, 4, 5 STRICT INTEGER)", Font(name="Segoe UI", size=10, bold=True, color="FFFFFF"), score_hdr_fill),
        (10, "ค่าน้ำหนัก (Weight %)", attr_header_font, attr_fill),
        (11, "คะแนนถ่วงน้ำหนัก (=Score * Weight)", attr_header_font, attr_fill),
        (12, "คำอธิบาย & หลักฐานอ้างอิงจากไฟล์ Intake (Evidence Citation)", attr_header_font, attr_fill),
        (13, "แผนลดขอบเขตงาน & มาตรการลดความเสี่ยง (De-scoping Action)", attr_header_font, attr_fill)
    ]

    for r_idx, label, fnt, fll in attributes:
        cell = ws.cell(row=r_idx, column=1, value=label)
        cell.font = fnt
        cell.fill = fll
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = thin_border

    ws.column_dimensions['A'].width = 38

    # ตรวจสอบว่าชุดประเมินมีมิติ SDG (+S) หรือไม่ เพื่อสลับระหว่างโมเดล TELOS (5 เสา) และ TELOS+S (6 เสา)
    has_sdg = any(
        any(k in str(it.get("pillar", "")).lower() for k in ["sdg", "ยั่งยืน", "สิ่งแวดล้อม"])
        for it in assessment_data
    )

    # นับจำนวนคำถามต่อมิติแบบ Canonical เพื่อใช้ใน Two-Tier Normalized Weighting Fallback
    pillar_counts = {}
    pillar_cols = {"T": [], "E": [], "L": [], "O": [], "S": [], "SDG": []}
    for it in assessment_data:
        c_key, _ = get_canonical_pillar(it.get("pillar", ""), has_sdg=has_sdg)
        pillar_counts[c_key] = pillar_counts.get(c_key, 0) + 1

    # เติมข้อมูลคำถามจาก AI Agents ลงในแนวคอลัมน์ (Columns B, C, D, ...)
    start_col = 2
    for idx, item in enumerate(assessment_data):
        col = start_col + idx
        col_letter = get_column_letter(col)
        ws.column_dimensions[col_letter].width = 36

        c_key, _ = get_canonical_pillar(item.get("pillar", ""), has_sdg=has_sdg)
        if c_key in pillar_cols:
            pillar_cols[c_key].append(col_letter)

        # Row 4: รหัส
        c4 = ws.cell(row=4, column=col, value=item.get("id", f"Q-{idx+1:02d}"))
        c4.font = q_id_font
        c4.fill = q_id_fill
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.border = thin_border

        # Row 5: มิติ
        c5 = ws.cell(row=5, column=col, value=item.get("pillar", "มิติการประเมิน"))
        c5.font = bold_font
        c5.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c5.border = thin_border

        # Row 6: คำถาม Stress-Test
        c6 = ws.cell(row=6, column=col, value=item.get("question", ""))
        c6.font = regular_font
        c6.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c6.border = thin_border

        # Row 7: เหตุผล
        c7 = ws.cell(row=7, column=col, value=item.get("rationale", ""))
        c7.font = regular_font
        c7.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c7.border = thin_border

        # Row 8: เกณฑ์ Rubric
        c8 = ws.cell(row=8, column=col, value=item.get("rubric", ""))
        c8.font = regular_font
        c8.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c8.border = thin_border

        # Row 9: คะแนน (Strict Integer: 1, 2, 3, 4, 5 เด็ดขาด)
        raw_val = item.get("score", 3)
        try:
            score_int = int(round(float(raw_val)))
        except (ValueError, TypeError):
            score_int = 3
        score_int = max(1, min(5, score_int))

        c9 = ws.cell(row=9, column=col, value=score_int)
        c9.font = score_font
        c9.fill = score_cell_fill
        c9.alignment = Alignment(horizontal="center", vertical="center")
        c9.number_format = '0'
        c9.border = thin_border

        # Row 10: น้ำหนัก (Two-Tier Normalized Weighting Architecture)
        weight_val = None
        if "weight" in item and item["weight"] is not None:
            try:
                raw_w = str(item["weight"]).replace('%', '').strip()
                parsed_w = float(raw_w)
                if parsed_w > 1.0:
                    parsed_w = parsed_w / 100.0  # แปลง 20 หรือ 6.67 ให้เป็น 0.20 หรือ 0.0667
                weight_val = parsed_w
            except (ValueError, TypeError):
                weight_val = None

        if weight_val is None:
            c_key, p_wt = get_canonical_pillar(item.get("pillar", ""), has_sdg=has_sdg)
            p_count = max(1, pillar_counts.get(c_key, 1))
            weight_val = p_wt / p_count

        c10 = ws.cell(row=10, column=col, value=weight_val)
        c10.font = bold_font
        c10.alignment = Alignment(horizontal="center", vertical="center")
        c10.number_format = '0.0%'
        c10.border = thin_border

        # Row 11: คะแนนถ่วงน้ำหนัก (สูตร Excel)
        c11 = ws.cell(row=11, column=col, value=f"={col_letter}9*{col_letter}10")
        c11.font = bold_font
        c11.alignment = Alignment(horizontal="center", vertical="center")
        c11.number_format = '0.00'
        c11.border = thin_border

        # Row 12: หลักฐานอ้างอิง
        c12 = ws.cell(row=12, column=col, value=item.get("evidence", ""))
        c12.font = regular_font
        c12.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c12.border = thin_border

        # Row 13: มาตรการลดความเสี่ยง
        c13 = ws.cell(row=13, column=col, value=item.get("descope", ""))
        c13.font = regular_font
        c13.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c13.border = thin_border

    # คอลัมน์สรุปภาพรวม (Summary Column)
    summary_col_idx = start_col + len(assessment_data)
    sum_letter = get_column_letter(summary_col_idx)
    ws.column_dimensions[sum_letter].width = 36

    first_col = get_column_letter(start_col)
    last_col = get_column_letter(summary_col_idx - 1)

    c_s4 = ws.cell(row=4, column=summary_col_idx, value="สรุปภาพรวม (SUMMARY)")
    c_s4.font = q_id_font
    c_s4.fill = q_id_fill
    c_s4.alignment = Alignment(horizontal="center", vertical="center")
    c_s4.border = thin_border

    c_s5 = ws.cell(row=5, column=summary_col_idx, value="ผลรวมทุกมิติ (All Pillars)")
    c_s5.font = bold_font
    c_s5.alignment = Alignment(horizontal="center", vertical="center")
    c_s5.border = thin_border

    ws.cell(row=6, column=summary_col_idx, value=f"ประเมินคำถามทั้งสิ้น {len(assessment_data)} ข้อ").font = regular_font
    ws.cell(row=6, column=summary_col_idx).alignment = Alignment(wrap_text=True)
    ws.cell(row=6, column=summary_col_idx).border = thin_border

    ws.cell(row=7, column=summary_col_idx, value="สังเคราะห์จาก AI Agents Group 2").font = regular_font
    ws.cell(row=7, column=summary_col_idx).alignment = Alignment(wrap_text=True)
    ws.cell(row=7, column=summary_col_idx).border = thin_border

    ws.cell(row=8, column=summary_col_idx, value="เกณฑ์รวม: >=80 ผ่านสูง, >=60 มีเงื่อนไข, <60 เสี่ยงสูง | หากมีข้อใดได้ 1 = VETOED ตกเกณฑ์ทันที").font = regular_font
    ws.cell(row=8, column=summary_col_idx).alignment = Alignment(wrap_text=True)
    ws.cell(row=8, column=summary_col_idx).border = thin_border

    # ตรวจสอบว่ามีคำถามใดได้คะแนน 1 หรือไม่ (Knockout / Fatal-Flaw Detection)
    has_fatal_flaw = any(
        (isinstance(it.get("score"), (int, float)) and int(round(it["score"])) == 1)
        or str(it.get("score", "")).strip() == "1"
        for it in assessment_data
    )

    # คะแนนรวมเต็ม 100 คำนวณแบบ Normalized กับผลรวมค่าน้ำหนักจริง (Dynamic Normalization)
    c_tot_score = ws.cell(row=9, column=summary_col_idx, value=f'=IF({sum_letter}10>0, ({sum_letter}11/{sum_letter}10)*20, 0)')
    if has_fatal_flaw:
        c_tot_score.font = Font(name="Segoe UI", size=16, bold=True, color="9C0006")
        c_tot_score.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    else:
        c_tot_score.font = Font(name="Segoe UI", size=16, bold=True, color="006100")
        c_tot_score.fill = highlight_fill
    c_tot_score.number_format = '0.0'
    c_tot_score.alignment = Alignment(horizontal="center", vertical="center")
    c_tot_score.border = thin_border

    # ค่าน้ำหนักรวม
    c_tot_wt = ws.cell(row=10, column=summary_col_idx, value=f"=SUM({first_col}10:{last_col}10)")
    c_tot_wt.font = bold_font
    c_tot_wt.number_format = '0.0%'
    c_tot_wt.alignment = Alignment(horizontal="center", vertical="center")
    c_tot_wt.border = thin_border

    # ผลรวมคะแนนถ่วงน้ำหนัก
    c_tot_weighted = ws.cell(row=11, column=summary_col_idx, value=f"=SUM({first_col}11:{last_col}11)")
    c_tot_weighted.font = bold_font
    c_tot_weighted.number_format = '0.00'
    c_tot_weighted.alignment = Alignment(horizontal="center", vertical="center")
    c_tot_weighted.border = thin_border

    # คำตัดสิน Verdict ตามกฎ Knockout / Fatal-Flaw Gating Rule
    # หาก MIN ของคะแนน = 1 ให้ขึ้น VETOED ตกเกณฑ์ข้อบังคับวิกฤตทันที
    verdict_c = ws.cell(row=12, column=summary_col_idx, value=f'=IF(MIN({first_col}9:{last_col}9)=1, "VETOED / ตกเกณฑ์ข้อบังคับวิกฤต (คะแนนระดับ 1)", IF({sum_letter}9>=80, "ผ่านเกณฑ์ระดับสูง (HIGHLY VIABLE)", IF({sum_letter}9>=60, "ผ่านแบบมีเงื่อนไข (CONDITIONALLY VIABLE)", "ความเสี่ยงสูง (HIGH RISK)")))')
    if has_fatal_flaw:
        verdict_c.font = Font(name="Segoe UI", size=11, bold=True, color="9C0006")
        verdict_c.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    else:
        verdict_c.font = verdict_font
        verdict_c.fill = highlight_fill
    verdict_c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    verdict_c.border = thin_border

    ws.cell(row=13, column=summary_col_idx, value="ปฏิบัติตามแผนลดขอบเขตงาน (De-scoping Advisory)").font = bold_font
    ws.cell(row=13, column=summary_col_idx).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=13, column=summary_col_idx).border = thin_border

    # =========================================================================
    # ตารางสรุปคะแนนแยกตามมิติ TELOS+S (Pillar Score Breakdown: Rows 15-23)
    # =========================================================================
    ws.cell(row=15, column=1, value="ตารางสรุปคะแนนแยกตามมิติ TELOS+S (Pillar Score Breakdown)").font = title_font

    breakdown_headers = [
        (1, "มิติการประเมิน (TELOS+S Pillar)"),
        (2, "สัดส่วนน้ำหนัก (Weight %)"),
        (3, "คะแนนเต็มมิติ (Max Points)"),
        (4, "คะแนนที่ได้รับจริง (Earned Points)"),
        (5, "คะแนนเฉลี่ย (Score 1-5)"),
        (6, "ระดับความพร้อม (Readiness Level)")
    ]
    for b_col, b_title in breakdown_headers:
        b_cell = ws.cell(row=16, column=b_col, value=b_title)
        b_cell.font = attr_header_font
        b_cell.fill = attr_fill
        b_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        b_cell.border = thin_border

    if not has_sdg:
        pillar_configs = [
            (17, "T", "T : ด้านเทคโนโลยีและวิศวกรรม (Technical)", 0.20, 20.0, 16.0, 12.0),
            (18, "E", "E : ด้านเศรษฐศาสตร์และความคุ้มค่า (Economic)", 0.20, 20.0, 16.0, 12.0),
            (19, "L", "L : ด้านกฎหมายและสถาบัน (Legal & Institutional)", 0.20, 20.0, 16.0, 12.0),
            (20, "O", "O : ด้านการปฏิบัติการและหน้างาน (Operational)", 0.20, 20.0, 16.0, 12.0),
            (21, "S", "S : ด้านแผนงานและเวลาส่งมอบ (Schedule)", 0.20, 20.0, 16.0, 12.0),
        ]
        tot_r = 22
        end_pillar_r = 21
    else:
        pillar_configs = [
            (17, "T", "T : ด้านเทคโนโลยีและวิศวกรรม (Technical)", 0.20, 20.0, 16.0, 12.0),
            (18, "E", "E : ด้านเศรษฐศาสตร์และความคุ้มค่า (Economic)", 0.20, 20.0, 16.0, 12.0),
            (19, "L", "L : ด้านกฎหมาย ระเบียบ และภาคประชาชน (Legal & Outreach)", 0.15, 15.0, 12.0, 9.0),
            (20, "O", "O : ด้านการปฏิบัติการและหน้างาน (Operational)", 0.20, 20.0, 16.0, 12.0),
            (21, "S", "S : ด้านแผนงานและเวลาส่งมอบ (Schedule)", 0.15, 15.0, 12.0, 9.0),
            (22, "SDG", "+S : ด้านความยั่งยืนและสิ่งแวดล้อม (Sustainability / SDG)", 0.10, 10.0, 8.0, 6.0),
        ]
        tot_r = 23
        end_pillar_r = 22

    for r_num, p_key, p_name, p_weight, p_max, th_high, th_med in pillar_configs:
        # Col A: ชื่อมิติ
        cA = ws.cell(row=r_num, column=1, value=p_name)
        cA.font = bold_font
        cA.border = thin_border

        # Col B: สัดส่วนน้ำหนัก
        cB = ws.cell(row=r_num, column=2, value=p_weight)
        cB.font = regular_font
        cB.number_format = '0.0%'
        cB.alignment = Alignment(horizontal="center", vertical="center")
        cB.border = thin_border

        # Col C: คะแนนเต็ม
        cC = ws.cell(row=r_num, column=3, value=p_max)
        cC.font = regular_font
        cC.number_format = '0.0'
        cC.alignment = Alignment(horizontal="center", vertical="center")
        cC.border = thin_border

        # Col D: คะแนนที่ได้รับจริง (สูตร Excel)
        p_formula = make_pillar_sum_formula(pillar_cols.get(p_key, []), 11)
        cD = ws.cell(row=r_num, column=4, value=p_formula)
        cD.font = Font(name="Segoe UI", size=10, bold=True, color="002060")
        cD.fill = score_cell_fill
        cD.number_format = '0.00'
        cD.alignment = Alignment(horizontal="center", vertical="center")
        cD.border = thin_border

        # Col E: คะแนนเฉลี่ย 1-5
        cE = ws.cell(row=r_num, column=5, value=f"=IF(C{r_num}>0, (D{r_num}/C{r_num})*5, 0)")
        cE.font = regular_font
        cE.number_format = '0.00'
        cE.alignment = Alignment(horizontal="center", vertical="center")
        cE.border = thin_border

        # Col F: สถานะความพร้อม
        cF = ws.cell(row=r_num, column=6, value=f'=IF(D{r_num}>={th_high}, "พร้อมระดับสูง (High)", IF(D{r_num}>={th_med}, "ปานกลาง (Medium)", "เสี่ยงวิกฤต (Critical)"))')
        cF.font = bold_font
        cF.alignment = Alignment(horizontal="center", vertical="center")
        cF.border = thin_border

    # แถวสรุปคะแนนรวมทั้งหมด
    cTotA = ws.cell(row=tot_r, column=1, value="คะแนนรวมทุกมิติ (Total Feasibility Score)")
    cTotA.font = bold_font
    cTotA.alignment = Alignment(horizontal="left", vertical="center")
    cTotA.fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    cTotA.border = thin_border

    cB_tot = ws.cell(row=tot_r, column=2, value=f"=SUM(B17:B{end_pillar_r})")
    cB_tot.font = bold_font
    cB_tot.number_format = '0.0%'
    cB_tot.alignment = Alignment(horizontal="center", vertical="center")
    cB_tot.border = thin_border

    cC_tot = ws.cell(row=tot_r, column=3, value=f"=SUM(C17:C{end_pillar_r})")
    cC_tot.font = bold_font
    cC_tot.number_format = '0.0'
    cC_tot.alignment = Alignment(horizontal="center", vertical="center")
    cC_tot.border = thin_border

    cD_tot = ws.cell(row=tot_r, column=4, value=f"=SUM(D17:D{end_pillar_r})")
    cD_tot.font = Font(name="Segoe UI", size=11, bold=True, color="002060")
    cD_tot.fill = highlight_fill
    cD_tot.number_format = '0.00'
    cD_tot.alignment = Alignment(horizontal="center", vertical="center")
    cD_tot.border = thin_border

    cE_tot = ws.cell(row=tot_r, column=5, value=f"=IF(C{tot_r}>0, (D{tot_r}/C{tot_r})*5, 0)")
    cE_tot.font = bold_font
    cE_tot.number_format = '0.00'
    cE_tot.alignment = Alignment(horizontal="center", vertical="center")
    cE_tot.border = thin_border

    cF_tot = ws.cell(row=tot_r, column=6, value=f'=IF(D{tot_r}>=80, "ผ่านเกณฑ์ระดับสูง (HIGHLY VIABLE)", IF(D{tot_r}>=60, "ผ่านแบบมีเงื่อนไข (CONDITIONALLY VIABLE)", "ความเสี่ยงสูง (HIGH RISK)"))')
    cF_tot.font = verdict_font
    cF_tot.fill = highlight_fill
    cF_tot.alignment = Alignment(horizontal="center", vertical="center")
    cF_tot.border = thin_border

    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 24
    ws.row_dimensions[6].height = 80
    ws.row_dimensions[7].height = 95
    ws.row_dimensions[8].height = 150
    ws.row_dimensions[9].height = 36
    ws.row_dimensions[10].height = 22
    ws.row_dimensions[11].height = 24
    ws.row_dimensions[12].height = 115
    ws.row_dimensions[13].height = 70
    ws.row_dimensions[14].height = 15
    ws.row_dimensions[15].height = 26
    ws.row_dimensions[16].height = 24
    for r in range(17, end_pillar_r + 1):
        ws.row_dimensions[r].height = 22
    ws.row_dimensions[tot_r].height = 26

    # =========================================================================
    # เมทริกซ์ประเมินความเข้ากันได้กับองค์ความรู้ด้านวิศวกรรมหุ่นยนต์ (Robotics Knowledge Compatibility)
    # เกณฑ์เฉพาะทาง FIBO (เต็ม 10 คะแนน) - เป็นเอกเทศ ไม่นับรวมในคะแนน TELOS 100 คะแนน
    # =========================================================================
    if robotic_data:
        r_spacer = tot_r + 1
        r_title = tot_r + 2
        r_note = tot_r + 3
        r_hdr = tot_r + 4
        r_data = tot_r + 5

        ws.row_dimensions[r_spacer].height = 12  # spacer
        ws.row_dimensions[r_title].height = 26
        ws.row_dimensions[r_note].height = 22
        ws.row_dimensions[r_hdr].height = 24
        ws.row_dimensions[r_data].height = 95

        ws.cell(row=r_title, column=1, value="เมทริกซ์ประเมินความเข้ากันได้กับองค์ความรู้ด้านวิศวกรรมหุ่นยนต์ (Robotics Knowledge Compatibility)").font = title_font
        ws.cell(row=r_note, column=1, value="* หมายเหตุ: การประเมินส่วนนี้วัดศักยภาพในการประยุกต์ใช้องค์ความรู้ด้านวิศวกรรมหุ่นยนต์ (Perception, Control, Actuation, Mechatronics) เต็ม 10 คะแนน | เป็นเอกเทศ ไม่นับรวมในคะแนน TELOS 100 คะแนน").font = sub_font

        robotic_headers = [
            (1, "มิติการประเมิน (Robotics Dimension)"),
            (2, "คะแนนศักยภาพ (/10)"),
            (3, "แกน Perception & Sensing (การรับรู้)"),
            (4, "แกน Control & Algorithms (การควบคุม)"),
            (5, "แกน Actuation & Mechanisms (กลไก)"),
            (6, "บทวิเคราะห์เชิงวิศวกรรมหุ่นยนต์ (FIBO Alignment Rationale)")
        ]
        r_hdr_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        for r_col, r_title_txt in robotic_headers:
            r_c = ws.cell(row=r_hdr, column=r_col, value=r_title_txt)
            r_c.font = attr_header_font
            r_c.fill = r_hdr_fill
            r_c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            r_c.border = thin_border

        r_data_A = ws.cell(row=r_data, column=1, value="ศักยภาพการประยุกต์ใช้องค์ความรู้หุ่นยนต์และระบบอัตโนมัติ")
        r_data_A.font = bold_font
        r_data_A.alignment = Alignment(horizontal="left", vertical="center")
        r_data_A.border = thin_border

        r_data_B = ws.cell(row=r_data, column=2, value=float(robotic_data.get("score", 0.0)))
        r_data_B.font = Font(name="Segoe UI", size=14, bold=True, color="002060")
        r_data_B.fill = score_cell_fill
        r_data_B.alignment = Alignment(horizontal="center", vertical="center")
        r_data_B.number_format = '0.0'
        r_data_B.border = thin_border

        domains = robotic_data.get("domains", {})
        r_data_C = ws.cell(row=r_data, column=3, value=domains.get("perception", "-"))
        r_data_C.font = regular_font
        r_data_C.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        r_data_C.border = thin_border

        r_data_D = ws.cell(row=r_data, column=4, value=domains.get("control_algorithms", "-"))
        r_data_D.font = regular_font
        r_data_D.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        r_data_D.border = thin_border

        r_data_E = ws.cell(row=r_data, column=5, value=domains.get("actuation_mechanics", "-"))
        r_data_E.font = regular_font
        r_data_E.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        r_data_E.border = thin_border

        r_data_F = ws.cell(row=r_data, column=6, value=robotic_data.get("fibo_alignment_rationale", "-"))
        r_data_F.font = regular_font
        r_data_F.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        r_data_F.border = thin_border

    return sum_letter



# ==============================================================================
# 2. ฟังก์ชันสร้างหน้าสรุปเปรียบเทียบทุก Solution (Executive Comparison Tab)
# ==============================================================================

def write_comparison_sheet(wb, solutions_info, has_sdg=False):
    ws = wb.create_sheet(title="เปรียบเทียบทุก Solution", index=0)
    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
    sub_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    hdr_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    bold_data_font = Font(name="Segoe UI", size=10, bold=True)
    score_sum_font = Font(name="Segoe UI", size=13, bold=True, color="006100")

    hdr_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    highlight_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    ws.cell(row=1, column=1, value="ตารางเปรียบเทียบความเป็นไปได้ของทุกข้อเสนอ (Multi-Solution Feasibility Summary)").font = title_font
    sub_title = ("วิเคราะห์เปรียบเทียบเชิงระบบตามเกณฑ์ TELOS+S (เต็ม 100) และเมทริกซ์ศักยภาพด้านวิศวกรรมหุ่นยนต์ (เต็ม 10 แยกอิสระ)" 
                 if has_sdg else 
                 "วิเคราะห์เปรียบเทียบเชิงระบบตามเกณฑ์ TELOS 5 เสาหลัก (เต็ม 100) และเมทริกซ์ศักยภาพด้านวิศวกรรมหุ่นยนต์ (เต็ม 10 แยกอิสระ)")
    ws.cell(row=2, column=1, value=sub_title).font = sub_font

    if not has_sdg:
        headers = [
            ("ลำดับ", 8),
            ("รหัสข้อเสนอ (Solution ID)", 22),
            ("ชื่อแนวทางแก้ปัญหา (Title / Concept)", 34),
            ("คะแนนรวม (เต็ม 100)", 18),
            ("T: เทคนิค (/20)", 14),
            ("E: เศรษฐศาสตร์ (/20)", 17),
            ("L: กฎหมายและสถาบัน (/20)", 18),
            ("O: ปฏิบัติการ (/20)", 16),
            ("S: แผนงาน (/20)", 14),
            ("ผลการประเมิน (Verdict)", 28),
            ("จุดแข็งสำคัญ (Key Strengths)", 35),
            ("จุดติดขัดวิกฤต (Critical Bottlenecks)", 35),
            ("คำแนะนำเชิงระบบ (Strategic Advice)", 35),
            ("Robotics Potential (/10)", 22),
            ("เหตุผลความเข้ากันได้ด้านหุ่นยนต์ (Robotics Engineering Rationale)", 45)
        ]
        robotic_cols = (14, 15)
        total_cols = 15
        verdict_col = 10
    else:
        headers = [
            ("ลำดับ", 8),
            ("รหัสข้อเสนอ (Solution ID)", 22),
            ("ชื่อแนวทางแก้ปัญหา (Title / Concept)", 34),
            ("คะแนนรวม (เต็ม 100)", 18),
            ("T: เทคนิค (/20)", 14),
            ("E: เศรษฐศาสตร์ (/20)", 17),
            ("L: กฎหมาย (/15)", 14),
            ("O: ปฏิบัติการ (/20)", 16),
            ("S: แผนงาน (/15)", 14),
            ("+S: ยั่งยืน (/10)", 14),
            ("ผลการประเมิน (Verdict)", 28),
            ("จุดแข็งสำคัญ (Key Strengths)", 35),
            ("จุดติดขัดวิกฤต (Critical Bottlenecks)", 35),
            ("คำแนะนำเชิงระบบ (Strategic Advice)", 35),
            ("Robotics Potential (/10)", 22),
            ("เหตุผลความเข้ากันได้ด้านหุ่นยนต์ (Robotics Engineering Rationale)", 45)
        ]
        robotic_cols = (15, 16)
        total_cols = 16
        verdict_col = 11

    robotic_hdr_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

    ws.row_dimensions[4].height = 28
    for col_idx, (hdr_text, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=hdr_text)
        cell.font = hdr_font
        cell.fill = robotic_hdr_fill if col_idx in robotic_cols else hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    for r_idx, sol in enumerate(solutions_info, start=5):
        ws.row_dimensions[r_idx].height = 55
        fill_to_use = zebra_fill if r_idx % 2 == 1 else PatternFill(fill_type=None)

        c1 = ws.cell(row=r_idx, column=1, value=r_idx - 4)
        c1.alignment = Alignment(horizontal="center", vertical="center")

        c2 = ws.cell(row=r_idx, column=2, value=sol["sheet_title"])
        c2.font = bold_data_font
        c2.alignment = Alignment(horizontal="center", vertical="center")

        c3 = ws.cell(row=r_idx, column=3, value=sol["concept"])
        c3.font = data_font
        c3.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        safe_sheet_name = sol['sheet_title'].replace("'", "''")
        sheet_ref = f"'{safe_sheet_name}'"
        c4 = ws.cell(row=r_idx, column=4, value=f"={sheet_ref}!{sol['sum_col']}9")
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.number_format = '0.0'

        # Pillar Breakdown (T, E, L, O, S)
        c5 = ws.cell(row=r_idx, column=5, value=f"={sheet_ref}!D17")
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.number_format = '0.00'
        c5.font = bold_data_font

        c6 = ws.cell(row=r_idx, column=6, value=f"={sheet_ref}!D18")
        c6.alignment = Alignment(horizontal="center", vertical="center")
        c6.number_format = '0.00'
        c6.font = bold_data_font

        c7 = ws.cell(row=r_idx, column=7, value=f"={sheet_ref}!D19")
        c7.alignment = Alignment(horizontal="center", vertical="center")
        c7.number_format = '0.00'
        c7.font = bold_data_font

        c8 = ws.cell(row=r_idx, column=8, value=f"={sheet_ref}!D20")
        c8.alignment = Alignment(horizontal="center", vertical="center")
        c8.number_format = '0.00'
        c8.font = bold_data_font

        c9 = ws.cell(row=r_idx, column=9, value=f"={sheet_ref}!D21")
        c9.alignment = Alignment(horizontal="center", vertical="center")
        c9.number_format = '0.00'
        c9.font = bold_data_font

        if has_sdg:
            c10 = ws.cell(row=r_idx, column=10, value=f"={sheet_ref}!D22")
            c10.alignment = Alignment(horizontal="center", vertical="center")
            c10.number_format = '0.00'
            c10.font = bold_data_font

        c_v = ws.cell(row=r_idx, column=verdict_col, value=f"={sheet_ref}!{sol['sum_col']}12")
        c_v.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        if sol.get("has_fatal_flaw", False):
            c4.font = Font(name="Segoe UI", size=13, bold=True, color="9C0006")
            c4.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            c_v.font = Font(name="Segoe UI", size=10, bold=True, color="9C0006")
            c_v.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        else:
            c4.font = score_sum_font
            c4.fill = highlight_fill
            c_v.font = bold_data_font

        next_c = verdict_col + 1
        ws.cell(row=r_idx, column=next_c, value=sol.get("strength", "-")).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=r_idx, column=next_c + 1, value=sol.get("bottleneck", "-")).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.cell(row=r_idx, column=next_c + 2, value=sol.get("advice", "-")).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        # เมทริกซ์ความเข้ากันได้ด้านหุ่นยนต์ (Robotics Knowledge Compatibility: แยกอิสระ เต็ม 10 คะแนน)
        rob_score_col = next_c + 3
        rob_rat_col = next_c + 4

        c_rob = ws.cell(row=r_idx, column=rob_score_col, value=float(sol.get("robotic_score", 0.0)))
        c_rob.font = Font(name="Segoe UI", size=11, bold=True, color="002060")
        c_rob.fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
        c_rob.alignment = Alignment(horizontal="center", vertical="center")
        c_rob.number_format = '0.0'

        c_rob_rat = ws.cell(row=r_idx, column=rob_rat_col, value=sol.get("robotic_rationale", "-"))
        c_rob_rat.font = data_font
        c_rob_rat.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

        for col_idx in range(1, total_cols + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = thin_border
            if col_idx not in (4, verdict_col, rob_score_col) and fill_to_use.fill_type:
                cell.fill = fill_to_use
            elif col_idx == verdict_col and not sol.get("has_fatal_flaw", False) and fill_to_use.fill_type:
                cell.fill = fill_to_use



# ==============================================================================
# 3. ฟังก์ชันบันทึกไฟล์ Excel ตามลำดับ _0, _1, _2
# ==============================================================================

def compile_excel_from_eval_data(solutions_eval_data, output_dir="feasibility-outcome", custom_filename=None):
    import re
    os.makedirs(output_dir, exist_ok=True)

    # ป้องกันข้อผิดพลาด Gap 1: กรณี LLM ห่อ List ด้วย Dictionary เช่น {"solutions": [...]}
    if isinstance(solutions_eval_data, dict):
        for candidate_key in ["solutions", "data", "evaluations", "results", "proposals"]:
            if candidate_key in solutions_eval_data and isinstance(solutions_eval_data[candidate_key], list):
                solutions_eval_data = solutions_eval_data[candidate_key]
                break
        else:
            for val in solutions_eval_data.values():
                if isinstance(val, list):
                    solutions_eval_data = val
                    break

    if not isinstance(solutions_eval_data, list):
        print("[ข้อผิดพลาด] รูปแบบข้อมูลการประเมินต้องเป็น List ของ Solutions!")
        return

    if not custom_filename:
        now = datetime.now()
        base_name = f"{now.month}-{now.day}-{now.year}-{now.strftime('%H%M')}"
        counter = 0
        filename = f"{base_name}_{counter}.xlsx"
        while os.path.exists(os.path.join(output_dir, filename)):
            counter += 1
            filename = f"{base_name}_{counter}.xlsx"
    elif not custom_filename.endswith(".xlsx"):
        filename = f"{custom_filename}.xlsx"
    else:
        filename = custom_filename

    excel_path = os.path.join(output_dir, filename)

    wb = openpyxl.Workbook()
    default_sheet = wb.active

    solutions_summary_info = []
    used_titles = set()

    for idx, sol in enumerate(solutions_eval_data, start=1):
        # ป้องกันข้อผิดพลาด Gap 3: ตัดอักขระต้องห้ามของ Excel \ / ? * : [ ] และป้องกันชื่อซ้ำ
        raw_title = str(sol.get("sheet_title", f"Solution-{idx}"))
        safe_title = re.sub(r'[\\/*?:\[\]]', '_', raw_title).strip()
        safe_title = safe_title[:28] if safe_title else f"Solution-{idx}"

        unique_title = safe_title
        dup_counter = 1
        while unique_title in used_titles or unique_title in wb.sheetnames:
            unique_title = f"{safe_title[:25]}_{dup_counter}"
            dup_counter += 1
        used_titles.add(unique_title)
        sheet_title = unique_title

        assessment_data = sol.get("assessment_data", [])
        concept = sol.get("concept", sheet_title)
        robotic_info = sol.get("robotic_compatibility") or sol.get("robotic_data") or {}
        robotic_score = robotic_info.get("score", 0.0)
        robotic_rationale = robotic_info.get("fibo_alignment_rationale", "-")

        sum_col = write_solution_sheet(wb, sheet_title, assessment_data, solution_name=concept, robotic_data=robotic_info)
        if not sum_col:
            continue

        has_fatal = any(
            (isinstance(it.get("score"), (int, float)) and int(round(it["score"])) == 1)
            or str(it.get("score", "")).strip() == "1"
            for it in assessment_data
        )

        solutions_summary_info.append({
            "sheet_title": sheet_title,
            "concept": concept,
            "sum_col": sum_col,
            "has_fatal_flaw": has_fatal,
            "strength": sol.get("strength", "-"),
            "bottleneck": sol.get("bottleneck", "-"),
            "advice": sol.get("advice", "-"),
            "robotic_score": robotic_score,
            "robotic_rationale": robotic_rationale
        })

    has_sdg_any = any(
        any(any(k in str(it.get("pillar", "")).lower() for k in ["sdg", "ยั่งยืน", "สิ่งแวดล้อม"])
            for it in sol.get("assessment_data", []))
        for sol in solutions_eval_data
    )

    write_comparison_sheet(wb, solutions_summary_info, has_sdg=has_sdg_any)

    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)

    try:
        wb.save(excel_path)
        print(f"[สำเร็จ] คอมไพล์เมทริกซ์ Excel จาก AI Agents เรียบร้อยแล้วที่:\n  --> {excel_path}")
        print(f"  --> จำนวนแท็บ Solution: {len(solutions_summary_info)} แท็บ + แท็บสรุปเปรียบเทียบ")
    except PermissionError:
        base, ext = os.path.splitext(excel_path)
        alt_path = f"{base}_alt{ext}"
        wb.save(alt_path)
        print(f"[แจ้งเตือน] ไฟล์ '{excel_path}' ถูกเปิดค้างอยู่ในโปรแกรมอื่น จึงบันทึกไปยัง:\n  --> {alt_path}")


def main():
    parser = argparse.ArgumentParser(description="คอมไพล์ผลการประเมิน TELOS+S จาก AI Agents เป็น Excel Matrix แนวนอน")
    parser.add_argument("--name", "-n", type=str, default=None, help="ชื่อไฟล์ Excel เช่น {month}-{date}-{year}-{time}")
    parser.add_argument("--input-json", "-i", type=str, default=None, help="ไฟล์ JSON ผลการประเมินจาก AI Agents")
    parser.add_argument("--outdir", "-o", type=str, default="feasibility-outcome", help="โฟลเดอร์สำหรับบันทึกไฟล์ผลลัพธ์")
    args = parser.parse_args()

    # ลำดับการหาไฟล์ input JSON จาก AI Agents
    json_path = args.input_json
    if not json_path:
        default_candidates = [
            os.path.join(args.outdir, "jury_eval_data.json"),
            "jury_eval_data.json",
            "eval_data.json"
        ]
        for candidate in default_candidates:
            if os.path.exists(candidate):
                json_path = candidate
                break

    if not json_path or not os.path.exists(json_path):
        print("[ข้อผิดพลาด] ไม่พบไฟล์ข้อมูลการประเมิน JSON จาก AI Agents!")
        print("กรุณาส่งพาธไฟล์ JSON ผ่าน: python feasibility-analysis.py --input-json <path_to_json>")
        sys.exit(1)

    with open(json_path, "r", encoding="utf-8") as f:
        solutions_eval_data = json.load(f)

    compile_excel_from_eval_data(solutions_eval_data, output_dir=args.outdir, custom_filename=args.name)


if __name__ == "__main__":
    main()
